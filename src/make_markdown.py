"""把两份 docx 导出成 Markdown，供 GitHub 网页直接渲染。

背景：docx 在 GitHub 上只能下载、不能预览，而这两份文档是本仓库的主要产出。
本脚本是 make_report_doc.py / make_defects_doc.py 的配套后处理，
不参与数据管线，改完 docx 后跑一次即可。

结构映射（依据两个生成脚本里各辅助函数的写法反推）：
    Title 样式              -> "# 标题"
    Heading N 样式          -> N+1 个 "#"（标题占 1 级，章节依次下移）
    List Bullet 样式        -> "- 列表项"，按左缩进还原嵌套层级
    Consolas 等宽字体段落    -> ``` 代码块，连续多段合并为一个
    Cambria Math 字体段落    -> "> *公式*" 引用行
    Table Grid 表格          -> Markdown 表格
    其余                    -> 普通段落，保留 **粗体** / *斜体*
"""
import os

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Pt
from docx.table import Table
from docx.text.paragraph import Paragraph

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))

DOCS = [
    ('复现报告.docx', '复现报告.md'),
    ('官方源码问题汇总.docx', '官方源码问题汇总.md'),
]

MONO_FONTS = {'Consolas', 'Courier New', 'Courier'}
MATH_FONTS = {'Cambria Math'}
BULLET_INDENT = Pt(18)          # bullet() 里每级缩进 18pt


def iter_blocks(doc):
    """按文档顺序产出段落和表格。

    doc.paragraphs 和 doc.tables 是两个独立列表，直接遍历会丢掉两者的相对顺序，
    所以这里走 XML body 的子元素。
    """
    for child in doc.element.body.iterchildren():
        if child.tag == qn('w:p'):
            yield Paragraph(child, doc)
        elif child.tag == qn('w:tbl'):
            yield Table(child, doc)


def classify(p):
    """判断段落类型，返回 (kind, level)。"""
    style = p.style.name or ''

    if style == 'Title':
        return 'title', 0
    if style.startswith('Heading'):
        try:
            return 'heading', int(style.split()[-1])
        except ValueError:
            return 'heading', 1
    if style.startswith('List Bullet'):
        return 'bullet', 0

    fonts = {r.font.name for r in p.runs if r.text.strip()}
    if fonts & MONO_FONTS:
        return 'code', 0
    if fonts & MATH_FONTS:
        return 'formula', 0
    return 'text', 0


def bullet_level(p):
    """从左缩进还原列表嵌套层级：18pt -> 0 级，36pt -> 1 级。"""
    indent = p.paragraph_format.left_indent
    if indent is None:
        return 0
    return max(0, round(indent / BULLET_INDENT) - 1)


def md_table(t):
    """Word 表格 -> Markdown 表格。单元格内的换行转 <br>，竖线转义。"""
    rows = []
    for row in t.rows:
        cells = []
        for c in row.cells:
            txt = ' '.join(p.text for p in c.paragraphs).strip()
            txt = txt.replace('|', r'\|').replace('\n', '<br>')
            cells.append(txt or ' ')
        rows.append(cells)

    if not rows:
        return ''

    width = max(len(r) for r in rows)
    out = ['| ' + ' | '.join(rows[0]) + ' |',
           '|' + '|'.join([' --- '] * width) + '|']
    for r in rows[1:]:
        r = r + [' '] * (width - len(r))
        out.append('| ' + ' | '.join(r) + ' |')
    return '\n'.join(out)


def render(doc, source_name):
    lines = ['<!-- 由 src/make_markdown.py 从 %s 自动导出，请勿直接编辑本文件 -->' % source_name,
             '',
             '> 本文件由 [`%s`](%s) 自动导出。需要 Word 版本请下载原文件。' % (source_name, source_name),
             '']

    code_buf = []

    def flush_code():
        if code_buf:
            lines.append('```')
            lines.extend(code_buf)
            lines.append('```')
            lines.append('')
            code_buf.clear()

    for blk in iter_blocks(doc):
        if isinstance(blk, Table):
            flush_code()
            md = md_table(blk)
            if md:
                lines.append(md)
                lines.append('')
            continue

        kind, level = classify(blk)
        text = blk.text.strip()

        if kind == 'code':
            if text:
                code_buf.extend(text.split('\n'))
            continue

        flush_code()

        if not text:
            continue

        if kind == 'title':
            lines.append('# ' + text)
        elif kind == 'heading':
            # 文档标题已占 1 级，章节从 2 级开始
            lines.append('#' * (level + 1) + ' ' + text)
        elif kind == 'bullet':
            lines.append('  ' * bullet_level(blk) + '- ' + text)
        elif kind == 'formula':
            lines.append('> *' + text.replace('\n', ' ') + '*')
        else:
            runs = [r for r in blk.runs if r.text.strip()]
            if runs and all(r.bold for r in runs):
                lines.append('**' + text + '**')
            elif runs and all(r.italic for r in runs):
                lines.append('*' + text.replace('\n', '  \n') + '*')
            else:
                lines.append(text.replace('\n', '  \n'))
        lines.append('')

    flush_code()

    # 折叠连续空行
    cleaned = []
    for ln in lines:
        if ln == '' and cleaned and cleaned[-1] == '':
            continue
        cleaned.append(ln)

    return '\n'.join(cleaned).rstrip() + '\n'


def main():
    for src, dst in DOCS:
        src_path = os.path.join(ROOT, src)
        dst_path = os.path.join(ROOT, dst)
        if not os.path.exists(src_path):
            print('跳过（源文件不存在）：', src_path)
            continue
        md = render(Document(src_path), src)
        with open(dst_path, 'w', encoding='utf-8', newline='\n') as f:
            f.write(md)
        print('已生成：%s  （%d 行，%d 字符）' % (dst_path, md.count('\n') + 1, len(md)))


if __name__ == '__main__':
    main()
